"""Excel import/export services and Qt workers for non-blocking file processing."""

import math
from pathlib import Path

from PySide6.QtCore import QObject, Signal, Slot

from customer_openpyxl_compat import ensure_openpyxl_numpy_compat


HEADER_MAP = {
    "區": "district", "地區": "district", "段": "section", "地段": "section",
    "序號": "registration_order", "登記次序": "registration_order",
    "地號": "land_number", "面積": "area", "面積/m2": "area", "面積m2": "area",
    "公告現值": "declared_value", "權利範圍": "rights_scope",
    "分子": "numerator", "分母": "denominator", "坪數": "ping",
    "總計公告現值": "total_declared_value", "總公告現值": "total_declared_value",
    "總現值/元": "total_declared_value", "總現值元": "total_declared_value",
    "所有權人": "owner_name", "姓名": "owner_name", "ID": "external_id",
    "身分證": "external_id", "身份證": "external_id", "地址": "address",
    "登記原因": "registration_reason", "原因": "registration_reason",
    "備註": "note", "出訪記錄": "visit_log",
}

HEADER_CONTAINS_MAP = [
    ("登記次序", "registration_order"), ("次序", "registration_order"),
    ("序號", "registration_order"), ("地號", "land_number"), ("面積", "area"),
    ("公告現值", "declared_value"), ("權利範圍", "rights_scope"),
    ("應有部分", "rights_scope"), ("持分", "rights_scope"),
    ("分子", "numerator"), ("分母", "denominator"), ("坪數", "ping"),
    ("總計公告現值", "total_declared_value"), ("總公告現值", "total_declared_value"),
    ("總現值", "total_declared_value"), ("所有權人", "owner_name"),
    ("所有人", "owner_name"), ("權利人", "owner_name"), ("姓名", "owner_name"),
    ("身分證", "external_id"), ("身份證", "external_id"),
    ("統一編號", "external_id"), ("統編", "external_id"), ("證號", "external_id"),
    ("地址", "address"), ("住址", "address"),
    ("登記原因", "registration_reason"), ("原因", "registration_reason"),
    ("備註", "note"), ("出訪", "visit_log"),
]

FORWARD_FILL_FIELDS = {
    "district", "section", "registration_order", "owner_name", "external_id",
    "address", "registration_reason", "rights_scope", "numerator", "denominator",
}

NUMERIC_IMPORT_FIELDS = {
    "area": "面積/m2", "declared_value": "公告現值", "numerator": "分子",
    "denominator": "分母", "ping": "坪數", "total_declared_value": "總現值/元",
}

REQUIRED_IMPORT_FIELDS = {
    "district": "地區",
    "section": "地段",
    "land_number": "地號",
}


def normalize_header(value):
    text = str(value or "").strip()
    for token in (" ", "\u3000", "\n", "\r", "\t", ":", "：", "(", ")", "（", "）"):
        text = text.replace(token, "")
    return text


def map_header(value):
    header = normalize_header(value)
    if not header:
        return None
    if header in HEADER_MAP:
        return HEADER_MAP[header]
    folded = header.casefold()
    if folded == "id":
        return "external_id"
    for token, key in HEADER_CONTAINS_MAP:
        if token in header:
            return key
    return None


def cell_to_text(value):
    if value is None:
        return None
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        if value.is_integer():
            return str(int(value))
    return str(value).strip()


def build_merged_value_map(sheet):
    merged_values = {}
    for merged_range in sheet.merged_cells.ranges:
        top_left_value = sheet.cell(merged_range.min_row, merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for column in range(merged_range.min_col, merged_range.max_col + 1):
                merged_values[(row, column)] = top_left_value
    return merged_values


def read_cell_value(sheet, row, column, merged_values):
    return merged_values.get((row, column), sheet.cell(row, column).value)


def find_header_row(sheet, merged_values, max_scan_rows=10):
    best_row = None
    best_column_map = {}
    scan_limit = min(max_scan_rows, sheet.max_row)
    for row_number in range(1, scan_limit + 1):
        column_map = {}
        mapped_keys = set()
        for column in range(1, sheet.max_column + 1):
            key = map_header(read_cell_value(sheet, row_number, column, merged_values))
            if key and key not in mapped_keys:
                column_map[column] = key
                mapped_keys.add(key)
        if len(column_map) > len(best_column_map):
            best_row = row_number
            best_column_map = column_map
    return best_row, best_column_map


class ExcelService:
    def __init__(
        self,
        land_fields,
        parse_number,
        split_rights_scope,
        format_number_text,
        calculate_ping,
        format_ping_text,
        calculate_total_declared_value,
    ):
        self.land_fields = tuple(land_fields)
        self.parse_number = parse_number
        self.split_rights_scope = split_rights_scope
        self.format_number_text = format_number_text
        self.calculate_ping = calculate_ping
        self.format_ping_text = format_ping_text
        self.calculate_total_declared_value = calculate_total_declared_value

    def validate_import_record(self, data):
        errors = []
        for key, label in NUMERIC_IMPORT_FIELDS.items():
            text = str(data.get(key) or "").strip()
            if text and self.parse_number(text) is None:
                errors.append(f"{label} 不是有效數字")
        denominator = self.parse_number(data.get("denominator"))
        if str(data.get("denominator") or "").strip() and denominator == 0:
            errors.append("分母不能為 0")
        if str(data.get("rights_scope") or "").strip():
            numerator, denominator_text = self.split_rights_scope(data.get("rights_scope"))
            if (
                not str(data.get("numerator") or "").strip()
                and not str(data.get("denominator") or "").strip()
                and numerator is None
                and denominator_text is None
            ):
                errors.append("權利範圍格式無法解析")
        for key, label in REQUIRED_IMPORT_FIELDS.items():
            if not str(data.get(key) or "").strip():
                errors.append(f"{label}不可空白")
        return errors

    def prepare_import_records(self, sheet, header_row, column_map, merged_values):
        records, error_rows, last_values = [], [], {}
        for row_number in range(header_row + 1, sheet.max_row + 1):
            data = {key: None for key, _label in self.land_fields}
            for column, key in column_map.items():
                data[key] = cell_to_text(read_cell_value(sheet, row_number, column, merged_values))
            if not any(data.values()):
                continue
            for key in FORWARD_FILL_FIELDS:
                if data.get(key):
                    last_values[key] = data[key]
                elif key in last_values:
                    data[key] = last_values[key]
            if data.get("rights_scope") and (not data.get("numerator") or not data.get("denominator")):
                numerator, denominator = self.split_rights_scope(data["rights_scope"])
                data["numerator"] = data.get("numerator") or numerator
                data["denominator"] = data.get("denominator") or denominator
            data["declared_value"] = self.format_number_text(data.get("declared_value") or "")
            data["ping"] = self.calculate_ping(data) or self.format_ping_text(data.get("ping") or "")
            data["total_declared_value"] = (
                self.calculate_total_declared_value(data)
                if not data.get("total_declared_value")
                else self.format_number_text(data["total_declared_value"])
            )
            data["name"] = data.get("owner_name") or ""
            errors = self.validate_import_record(data)
            if errors:
                error_rows.append(
                    {"row_number": row_number, "reason": "、".join(errors), "data": dict(data)}
                )
            else:
                records.append(data)
        return records, error_rows

    def read_import_file(self, file_path):
        try:
            ensure_openpyxl_numpy_compat()
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("目前環境缺少 openpyxl，無法處理 Excel 匯入。") from exc
        workbook = load_workbook(file_path, read_only=False, data_only=True)
        try:
            sheet = workbook.active
            merged_values = build_merged_value_map(sheet)
            header_row, column_map = find_header_row(sheet, merged_values)
            if not header_row or not column_map:
                raise ValueError("找不到可辨識的 Excel 欄位標題。")
            records, error_rows = self.prepare_import_records(
                sheet, header_row, column_map, merged_values
            )
            return {
                "source_file_name": Path(file_path).name,
                "column_map": column_map,
                "records": records,
                "error_rows": error_rows,
            }
        finally:
            workbook.close()

    @staticmethod
    def write_rows(file_path, rows, columns, sheet_title="客戶資料"):
        try:
            ensure_openpyxl_numpy_compat()
            from openpyxl import Workbook
        except ImportError as exc:
            raise RuntimeError("目前環境缺少 openpyxl，無法匯出 Excel。") from exc
        workbook = Workbook()
        try:
            sheet = workbook.active
            sheet.title = sheet_title
            for column_index, (_key, label) in enumerate(columns, start=1):
                sheet.cell(row=1, column=column_index, value=label)
            for row_index, row in enumerate(rows, start=2):
                for column_index, (key, _label) in enumerate(columns, start=1):
                    sheet.cell(row=row_index, column=column_index, value=row.get(key, ""))
            workbook.save(file_path)
        finally:
            workbook.close()

    @staticmethod
    def write_error_rows(file_path, error_rows, land_fields):
        columns = [("row_number", "Excel 列號"), ("reason", "錯誤原因"), *land_fields]
        flattened_rows = []
        for error_row in error_rows:
            flattened = {
                "row_number": error_row.get("row_number"),
                "reason": error_row.get("reason"),
            }
            flattened.update(error_row.get("data", {}))
            flattened_rows.append(flattened)
        ExcelService.write_rows(file_path, flattened_rows, columns, "匯入錯誤列")


class ExcelImportWorker(QObject):
    finished = Signal(object)
    failed = Signal(str)

    def __init__(self, service, file_path):
        super().__init__()
        self.service = service
        self.file_path = file_path

    @Slot()
    def run(self):
        try:
            self.finished.emit(self.service.read_import_file(self.file_path))
        except Exception as exc:
            self.failed.emit(str(exc))


class ExcelExportWorker(QObject):
    finished = Signal(str, int)
    failed = Signal(str)

    def __init__(self, file_path, rows, columns):
        super().__init__()
        self.file_path = file_path
        self.rows = rows
        self.columns = columns

    @Slot()
    def run(self):
        try:
            ExcelService.write_rows(self.file_path, self.rows, self.columns)
            self.finished.emit(str(Path(self.file_path)), len(self.rows))
        except Exception as exc:
            self.failed.emit(str(exc))
