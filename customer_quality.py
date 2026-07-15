"""Data quality inspection rules for customer land records."""

from collections import defaultdict
from dataclasses import dataclass
import hashlib
import json

from customer_domain import normalize_match_text, parse_number
from customer_openpyxl_compat import ensure_openpyxl_numpy_compat


@dataclass(frozen=True)
class DataQualityIssue:
    record_id: int
    severity: str
    category: str
    summary: str
    suggestion: str


QUALITY_RULE_OPTIONS = (
    ("owner_name_missing", "姓名空白"),
    ("land_number_missing", "地號缺漏"),
    ("numeric_format", "數字格式異常（面積、公告現值、分子、分母）"),
    ("denominator_zero", "分母為 0"),
    ("duplicate", "疑似重複"),
)
DEFAULT_QUALITY_RULE_KEYS = tuple(key for key, _label in QUALITY_RULE_OPTIONS)
QUALITY_RULE_PRESETS = (
    (
        "complete",
        "完整檢查",
        DEFAULT_QUALITY_RULE_KEYS,
        "檢查全部規則，適合定期總整理。",
    ),
    (
        "quick",
        "快速檢查",
        (
            "owner_name_missing",
            "land_number_missing",
            "numeric_format",
            "denominator_zero",
        ),
        "先排除缺漏與格式問題，不檢查疑似重複。",
    ),
    (
        "missing",
        "只查缺漏",
        ("owner_name_missing", "land_number_missing"),
        "只檢查姓名與地號是否缺漏。",
    ),
    (
        "numeric",
        "只查數字",
        ("numeric_format", "denominator_zero"),
        "只檢查面積、公告現值、分子、分母與分母為 0。",
    ),
    (
        "duplicate",
        "只查重複",
        ("duplicate",),
        "只檢查同地號、同姓名的疑似重複資料。",
    ),
)


def quality_issue_signature(issue):
    payload = {
        "record_id": int(issue.record_id or 0),
        "severity": str(issue.severity or ""),
        "category": str(issue.category or ""),
        "summary": str(issue.summary or ""),
        "suggestion": str(issue.suggestion or ""),
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _text(value):
    return str(value or "").strip()


def _record_label(record):
    parts = [
        _text(record.get("district")),
        _text(record.get("section")),
        _text(record.get("registration_order")),
        _text(record.get("land_number")),
    ]
    label = " / ".join(part for part in parts if part)
    return label or "未填土地識別"


def _add_issue(issues, record, severity, category, detail, suggestion):
    record_id = int(record.get("id") or 0)
    issues.append(
        DataQualityIssue(
            record_id=record_id,
            severity=severity,
            category=category,
            summary=f"#{record_id} {_record_label(record)}：{detail}",
            suggestion=suggestion,
        )
    )


def inspect_customer_quality(records, enabled_rules=None):
    records = [dict(record) for record in records]
    enabled = set(DEFAULT_QUALITY_RULE_KEYS if enabled_rules is None else enabled_rules)
    issues = []
    duplicate_groups = defaultdict(list)

    for record in records:
        if "owner_name_missing" in enabled and not _text(record.get("owner_name")):
            _add_issue(
                issues,
                record,
                "提醒",
                "姓名空白",
                "姓名尚未填寫",
                "補上所有權人姓名；若確定未知，可在備註註明原因。",
            )

        if "land_number_missing" in enabled and not _text(record.get("land_number")):
            _add_issue(
                issues,
                record,
                "重要",
                "地號缺漏",
                "地號尚未填寫",
                "補上地號，避免搜尋、排序與匯出時無法辨識土地。",
            )

        if "numeric_format" in enabled:
            for key, label in (("area", "面積"), ("declared_value", "公告現值")):
                value = _text(record.get(key))
                if value and parse_number(value) is None:
                    _add_issue(
                        issues,
                        record,
                        "重要",
                        f"{label}格式異常",
                        f"{label}「{value}」不是可辨識的數字",
                        f"請改成數字格式，例如 120 或 20,000。",
                    )

        for key, label in (("numerator", "分子"), ("denominator", "分母")):
            value = _text(record.get(key))
            number = parse_number(value)
            if "numeric_format" in enabled and value and number is None:
                _add_issue(
                    issues,
                    record,
                    "重要",
                    f"{label}格式異常",
                    f"{label}「{value}」不是可辨識的數字",
                    "請輸入數字；若權利範圍未知，可先留空並在備註說明。",
                )
            elif "denominator_zero" in enabled and key == "denominator" and number == 0:
                _add_issue(
                    issues,
                    record,
                    "重要",
                    "分母為 0",
                    "分母不可為 0",
                    "請修正分母；分母為 0 會造成坪數與總現值無法正確計算。",
                )

        if "duplicate" in enabled:
            duplicate_key = (
                normalize_match_text(record.get("district")),
                normalize_match_text(record.get("section")),
                normalize_match_text(record.get("registration_order")),
                normalize_match_text(record.get("land_number")),
                normalize_match_text(record.get("owner_name")),
                normalize_match_text(record.get("external_id")),
            )
            if any(duplicate_key):
                duplicate_groups[duplicate_key].append(record)

    for group in duplicate_groups.values():
        if len(group) < 2:
            continue
        record_ids = ", ".join(f"#{int(record.get('id') or 0)}" for record in group)
        for record in group:
            _add_issue(
                issues,
                record,
                "提醒",
                "疑似重複",
                f"與 {record_ids} 有相同土地識別與所有權人資料",
                "請確認是否重複輸入；若確實為不同資料，可在備註補充差異。",
            )

    return sorted(issues, key=lambda issue: (issue.record_id, issue.category, issue.summary))


def write_quality_issues(file_path, issues):
    try:
        ensure_openpyxl_numpy_compat()
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("目前環境缺少 openpyxl，無法匯出 Excel。") from exc

    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = "資料品質檢查"
        columns = [
            ("record_id", "資料ID"),
            ("severity", "程度"),
            ("category", "類型"),
            ("summary", "問題"),
            ("suggestion", "建議處理"),
        ]
        for column_index, (_key, label) in enumerate(columns, start=1):
            sheet.cell(row=1, column=column_index, value=label)

        for row_index, issue in enumerate(issues, start=2):
            values = {
                "record_id": issue.record_id,
                "severity": issue.severity,
                "category": issue.category,
                "summary": issue.summary,
                "suggestion": issue.suggestion,
            }
            for column_index, (key, _label) in enumerate(columns, start=1):
                sheet.cell(row=row_index, column=column_index, value=values[key])

        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 10
        sheet.column_dimensions["B"].width = 10
        sheet.column_dimensions["C"].width = 18
        sheet.column_dimensions["D"].width = 60
        sheet.column_dimensions["E"].width = 60
        workbook.save(file_path)
    finally:
        workbook.close()


def write_quality_summary(file_path, summary_rows):
    try:
        ensure_openpyxl_numpy_compat()
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("目前環境缺少 openpyxl，無法匯出 Excel。") from exc

    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = "品質檢查摘要"
        columns = [
            ("category", "類型"),
            ("important", "重要"),
            ("reminder", "提醒"),
            ("total", "合計"),
        ]
        for column_index, (_key, label) in enumerate(columns, start=1):
            sheet.cell(row=1, column=column_index, value=label)

        for row_index, row in enumerate(summary_rows, start=2):
            values = {
                "category": row.get("category", ""),
                "important": row.get("important", 0),
                "reminder": row.get("reminder", 0),
                "total": row.get("total", 0),
            }
            for column_index, (key, _label) in enumerate(columns, start=1):
                sheet.cell(row=row_index, column=column_index, value=values[key])

        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 10
        sheet.column_dimensions["C"].width = 10
        sheet.column_dimensions["D"].width = 10
        workbook.save(file_path)
    finally:
        workbook.close()
