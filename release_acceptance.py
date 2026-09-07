"""Isolated release acceptance workflow runnable from Python or the packaged EXE."""

import json
import os
import shutil
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

from build_data_guard import copy_sqlite_snapshot, restore as restore_build_data, save as save_build_data
from customer_database import CustomerDatabase
from customer_backup_status import BackupManagementDialog
from customer_domain import (
    calculate_ping,
    calculate_total_declared_value,
    format_number_text,
    format_ping_text,
    parse_number,
    parse_shared_land_rows,
    split_rights_scope,
)
from customer_excel import ExcelService
from customer_repository import CustomerRepository
from customer_productivity import (
    ProductivityService,
    apply_default_import_profile,
    export_report_with_template,
)
from customer_search import CustomerDecryptionCache, CustomerRecordProcessor
from customer_security import decrypt_value, encrypt_record, make_fernet
from customer_version import APP_VERSION, BUILD_DATE
from customer_word import write_records_docx

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from PySide6.QtWidgets import QApplication


LAND_FIELDS = [
    ("district", "地區"),
    ("section", "地段"),
    ("subsection", "小段"),
    ("registration_order", "序號"),
    ("land_number", "地號"),
    ("area", "面積/m2"),
    ("declared_value", "公告現值"),
    ("numerator", "分子"),
    ("denominator", "分母"),
    ("ping", "坪數"),
    ("total_declared_value", "總現值/元"),
    ("owner_name", "姓名"),
    ("external_id", "身分證"),
    ("address", "地址"),
    ("registration_reason", "原因"),
    ("note", "備註"),
    ("visit_log", "出訪記錄"),
]


def resource_dir():
    if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"):
        return Path(sys._MEIPASS)
    return Path(__file__).resolve().parent


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def run_release_acceptance(report_path=None):
    report = {
        "success": False,
        "version": APP_VERSION,
        "build_date": BUILD_DATE,
        "executed_at": datetime.now().isoformat(timespec="seconds"),
        "steps": [],
    }

    def passed(name, detail=""):
        report["steps"].append({"name": name, "status": "passed", "detail": detail})

    try:
        application = QApplication.instance() or QApplication([])
        backup_dialog = BackupManagementDialog(
            compress_backups=True,
            retention_days=90,
            max_count=30,
        )
        require(
            backup_dialog.selected_policy()
            == {"compress_backups": True, "retention_days": 90, "max_count": 30},
            "backup management dialog policy contract failed",
        )
        backup_dialog.close()
        passed("ui_contracts", "Qt 與備份管理真實視窗建構正常")

        with tempfile.TemporaryDirectory(prefix="land-customer-release-") as temp_name:
            root = Path(temp_name)
            resources = resource_dir()
            database = CustomerDatabase(root / "customers.db", root / "backups")
            repository = CustomerRepository(
                database,
                resources / "schema.sql",
                resources / "seed.sql",
                LAND_FIELDS,
            )

            repository.init_db()
            database.validate_database_file(database.database_path)
            require(not repository.has_admin_user(), "fresh database unexpectedly has admin")
            initial_count = repository.count_customers()
            passed(
                "fresh_start",
                f"全新資料庫建立、{initial_count} 筆初始資料與完整性檢查正常",
            )

            password = "release-test-password"
            repository.create_admin_user(password)
            encryption_key = repository.authenticate_user("admin", password)
            require(encryption_key is not None, "admin authentication failed")
            fernet = make_fernet(encryption_key)
            passed("authentication", "首次設定密碼與重新登入正常")

            first = {
                "district": "桃園區",
                "section": "一段",
                "registration_order": "1",
                "land_number": "100",
                "area": "100",
                "declared_value": "20000",
                "numerator": "1",
                "denominator": "2",
                "ping": "15.12",
                "total_declared_value": "1000000",
                "owner_name": "驗收人員",
                "external_id": "H100000001",
                "address": "桃園市驗收路1號",
                "registration_reason": "驗收",
                "note": "release smoke",
                "visit_log": "初訪",
                "name": "驗收人員",
            }
            first_id = repository.save_customer(encrypt_record(fernet, first))
            require(first_id > 0, "failed to insert first record")

            parsed_rows, errors = parse_shared_land_rows(
                "15、王弘益、H100059743、桃園市桃園區大華九街34號、108、3360"
            )
            require(not errors and len(parsed_rows) == 1, "batch parser failed")
            batch = {
                **{key: None for key, _label in LAND_FIELDS},
                **parsed_rows[0],
                "district": "桃園區",
                "section": "二段",
                "land_number": "200",
                "area": "108",
                "declared_value": "30000",
                "name": parsed_rows[0]["owner_name"],
            }
            batch["ping"] = calculate_ping(batch)
            batch["total_declared_value"] = calculate_total_declared_value(batch)
            repository.insert_customers([encrypt_record(fernet, batch)])
            expected_count = initial_count + 2
            require(repository.count_customers() == expected_count, "batch insert count mismatch")
            passed("crud_and_batch", "新增與同地號批量資料正常")

            candidates = repository.fetch_search_candidate_rows(
                keyword="桃園區",
                filter_field="district",
            )
            cache = CustomerDecryptionCache()
            cache.sync_revision(repository.data_revision)
            processor = CustomerRecordProcessor(
                fernet=fernet,
                table_columns=[("checked", ""), *LAND_FIELDS],
                keyword="桃園區",
                filter_field="district",
                decryption_cache=cache,
            )
            search_rows = processor.process(candidates)
            require(len(search_rows) == 2, "search result count mismatch")
            passed("search", "SQLite 候選篩選與解密搜尋正常")

            export_path = root / "acceptance-export.xlsx"
            ExcelService.write_rows(
                export_path,
                [record["raw"] for record in search_rows],
                [("land_number", "地號"), ("owner_name", "姓名")],
            )
            from openpyxl import load_workbook

            workbook = load_workbook(export_path, read_only=True, data_only=True)
            try:
                require(workbook.active.max_row == 3, "Excel export row count mismatch")
            finally:
                workbook.close()
            passed("excel_export", "Excel 匯出檔可重新讀取")

            word_path = root / "acceptance-export.docx"
            write_records_docx(
                word_path,
                [record["raw"] for record in search_rows],
                [("land_number", "地號"), ("owner_name", "姓名")],
            )
            require(word_path.is_file() and word_path.stat().st_size > 500, "Word export failed")
            passed("word_export", "Word 選取資料匯出可正常產生 docx")

            case_id = repository.save_case("acceptance case")
            repository.add_customers_to_case(case_id, [first_id])
            tag_id = repository.save_tag("acceptance tag", "#ffcc00")
            repository.set_customer_tags(first_id, [tag_id])
            field_id = repository.save_custom_field("Acceptance Field", "acceptance_field")
            repository.set_customer_custom_values(first_id, {field_id: "custom value"})
            template_id = repository.save_text_template("Acceptance Template", "template body", "note")
            repository.add_customer_attachment(first_id, str(word_path), "word export")
            repository.add_contact_log(first_id, "2026-07-13", "phone", "ok", "2026-07-20", "contact note")
            repository.save_follow_up_reminder(first_id, "2026-07-20", "待回覆", "follow up")
            managed_row = repository.get_customer(first_id)
            require(managed_row["case_names"] == "acceptance case", "case metadata not visible")
            require(managed_row["tag_names"] == "acceptance tag", "tag metadata not visible")
            require(managed_row["primary_tag_color"] == "#ffcc00", "tag color not visible")
            require(managed_row["attachment_count"] == 1, "attachment count not visible")
            require("word export" in managed_row["attachment_names"], "attachment metadata not visible")
            require("Acceptance Field：custom value" in managed_row["custom_values"], "custom metadata not visible")
            require("phone" in managed_row["last_contact"], "contact metadata not visible")
            require(managed_row["next_follow_up"] == "2026-07-20", "follow-up date not visible")
            require(managed_row["follow_up_status"] == "待回覆", "follow-up status not visible")
            second_id = max(repository.fetch_customer_ids() - {first_id})
            merged = dict(first)
            merged["external_id"] = batch["external_id"]
            merged["note"] = "release smoke\nmerged"
            repository.merge_customers(first_id, second_id, encrypt_record(fernet, merged))
            expected_count -= 1
            require(repository.count_customers() == expected_count, "merge count mismatch")
            require(repository.get_customer_tag_ids(first_id) == {tag_id}, "tag assignment failed")
            require(repository.get_customer_custom_values(first_id)[field_id] == "custom value", "custom value failed")
            require(repository.list_text_templates()[0]["id"] == template_id, "template save failed")
            require(repository.list_customer_attachments(first_id), "attachment save failed")
            require(repository.list_contact_logs(first_id), "contact log save failed")
            passed("management_features", "管理欄位顯示/標籤顏色/案件/附件/自訂欄位/聯絡/提醒/合併流程通過")

            editor_id = repository.create_user(
                "acceptance-editor",
                "acceptance-editor-password",
                "editor",
                encryption_key,
                "驗收編輯者",
            )
            editor_key = repository.authenticate_user(
                "acceptance-editor", "acceptance-editor-password"
            )
            require(editor_id > 0 and editor_key == encryption_key, "shared-key user login failed")

            today_text = date.today().isoformat()
            repository.save_case(
                "acceptance case",
                "進行中",
                "workflow",
                case_id,
                assigned_to="acceptance-editor",
                due_date=today_text,
                priority="高",
                next_action="release check",
            )
            task_id = repository.save_case_task(
                case_id,
                "verify release",
                assignee="acceptance-editor",
                due_date=today_text,
                priority="緊急",
            )
            repository.refresh_notifications(today_text)
            require(task_id > 0 and repository.list_notifications(), "workflow notifications failed")

            repository.save_import_profile(
                "acceptance profile",
                {"驗收所有權人": "owner_name"},
                is_default=True,
            )
            require(apply_default_import_profile(repository) == 1, "import profile failed")
            report_template_id = repository.save_report_template(
                "acceptance report",
                "驗收地主清冊",
                ["district", "land_number", "owner_name"],
                header_text="release acceptance",
            )
            report_template = dict(repository.list_report_templates()[0])
            templated_word = root / "acceptance-template.docx"
            export_report_with_template(
                templated_word,
                [first],
                report_template,
                dict(LAND_FIELDS),
            )
            require(
                report_template_id > 0 and templated_word.stat().st_size > 500,
                "report template export failed",
            )

            managed_source = root / "managed-attachment.txt"
            managed_source.write_text("managed release attachment", encoding="utf-8")
            managed_attachment_id = repository.import_managed_attachment(
                first_id, managed_source, root / "attachments", "managed"
            )
            attachment_results = repository.verify_managed_attachments()
            require(
                any(
                    item["id"] == managed_attachment_id and item["state"] == "正常"
                    for item in attachment_results
                ),
                "managed attachment verification failed",
            )

            productivity = ProductivityService(
                repository,
                database,
                root,
                root / "attachments",
                fernet,
            )
            mirror_directory = root / "offsite"
            repository.save_backup_target("acceptance target", mirror_directory)
            portable_backup, mirror_results = productivity.sync_backup_targets(
                "release-backup-password"
            )
            require(
                portable_backup.suffix == ".lcsbak"
                and productivity.verify_full_backup(
                    portable_backup, "release-backup-password"
                )
                and mirror_results[0][1] == "成功",
                "encrypted offsite full backup failed",
            )

            repository.set_customer_location(first_id, 25.033, 121.5654)
            map_path = root / "acceptance-map.html"
            productivity.build_map_html(repository.list_customer_locations(), map_path)
            require(map_path.is_file() and "OpenStreetMap" in map_path.read_text(encoding="utf-8"), "map export failed")

            duplicate = {**first, "name": first["owner_name"]}
            repository.insert_customers([encrypt_record(fernet, duplicate)])
            duplicate_id = repository.last_inserted_customer_ids[0]
            duplicate_pairs = productivity.find_duplicate_pairs(
                [{**first, "id": first_id}, {**first, "id": duplicate_id}]
            )
            require(duplicate_pairs and duplicate_pairs[0]["score"] == 100, "duplicate scoring failed")
            repository.record_insert_undo(
                "acceptance insert", [duplicate_id], "undo acceptance duplicate"
            )
            require(repository.undo_operation(), "batch undo failed")
            require(repository.get_customer(duplicate_id) is None, "batch undo did not remove insert")

            repository.delete_customers([first_id])
            recycle_item = repository.list_recycle_bin()[0]
            require(repository.get_customer(first_id) is None, "recycle delete failed")
            require(
                repository.restore_recycle_items([recycle_item["id"]]) == [first_id],
                "recycle restore failed",
            )
            passed(
                "productivity_suite",
                "回收桶/復原/多使用者/附件納管/異地加密備份/工作流/通知/匯入設定/報表/重複檢查/地圖通過",
            )

            repository.set_settings(
                {"checked_record_ids": f"[{first_id}]", "selected_record_id": str(first_id)}
            )
            reopened = CustomerRepository(
                database,
                resources / "schema.sql",
                resources / "seed.sql",
                LAND_FIELDS,
            )
            require(reopened.authenticate_user("admin", password), "reopen login failed")
            require(reopened.count_customers() == expected_count, "reopen record count mismatch")
            require(reopened.get_setting("selected_record_id") == str(first_id), "state not persisted")
            passed("reopen_state", "關閉重開所需資料與選取狀態正常保存")

            backup_path = database.backup_database("acceptance", compress=True)
            require(backup_path.suffix == ".zip", "compressed backup was not created")
            database.validate_backup_file(backup_path)
            repository.delete_all_customers()
            require(repository.count_customers() == 0, "pre-restore mutation failed")
            database.restore_database(backup_path)
            require(repository.count_customers() == expected_count, "backup restore count mismatch")
            passed("backup_restore", "ZIP 壓縮備份、驗證與直接還原正常")

            corrupt_path = root / "corrupt.db"
            corrupt_path.write_bytes(b"not a sqlite database")
            try:
                database.restore_database(corrupt_path)
            except ValueError:
                pass
            else:
                raise AssertionError("corrupt database restore was not rejected")
            require(
                repository.count_customers() == expected_count,
                "corrupt restore changed current data",
            )
            passed("corrupt_restore_guard", "損壞資料庫遭拒且現有資料未變更")

            build_root = root / "build-guard"
            deployed = build_root / "dist" / "LandCustomerSystem"
            deployed.mkdir(parents=True)
            copy_sqlite_snapshot(database.database_path, deployed / "customers.db")
            deployed_backups = deployed / "backups"
            deployed_backups.mkdir()
            shutil.copy2(backup_path, deployed_backups / backup_path.name)
            deployed_logs = deployed / "logs"
            deployed_logs.mkdir()
            (deployed_logs / "application-error.log").write_text(
                "release log marker",
                encoding="utf-8",
            )
            deployed_attachments = deployed / "attachments" / "7"
            deployed_attachments.mkdir(parents=True)
            (deployed_attachments / "managed.txt").write_text(
                "release attachment marker", encoding="utf-8"
            )
            save_build_data(build_root)
            shutil.rmtree(build_root / "dist")
            restore_build_data(build_root)
            guarded_db = deployed / "customers.db"
            database.validate_database_file(guarded_db)
            require((deployed_backups / backup_path.name).is_file(), "backup files not restored")
            require(
                (deployed_logs / "application-error.log").read_text(encoding="utf-8")
                == "release log marker",
                "application logs not restored",
            )
            require(
                (deployed_attachments / "managed.txt").read_text(encoding="utf-8")
                == "release attachment marker",
                "managed attachments not restored",
            )
            passed("rebuild_data_guard", "重新打包資料庫、備份、納管附件與錯誤記錄保存正常")

        report["success"] = True
    except Exception as exc:
        report["error"] = f"{type(exc).__name__}: {exc}"

    if report_path:
        report_path = Path(report_path)
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def main(report_path=None):
    return 0 if run_release_acceptance(report_path).get("success") else 1
