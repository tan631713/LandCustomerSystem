import unittest
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from customer_quality import (
    DataQualityIssue,
    inspect_customer_quality,
    quality_issue_signature,
    write_quality_summary,
    write_quality_issues,
)


class CustomerQualityTests(unittest.TestCase):
    def test_detects_missing_values_invalid_numbers_and_zero_denominator(self):
        issues = inspect_customer_quality(
            [
                {
                    "id": 1,
                    "district": "桃園市",
                    "section": "一段",
                    "registration_order": "15",
                    "land_number": "",
                    "owner_name": "",
                    "external_id": "",
                    "area": "abc",
                    "declared_value": "20,000",
                    "numerator": "1",
                    "denominator": "0",
                }
            ]
        )

        categories = {issue.category for issue in issues}
        self.assertIn("姓名空白", categories)
        self.assertIn("地號缺漏", categories)
        self.assertIn("面積格式異常", categories)
        self.assertIn("分母為 0", categories)

    def test_enabled_rules_limit_quality_issues(self):
        issues = inspect_customer_quality(
            [
                {
                    "id": 1,
                    "district": "桃園市",
                    "section": "一段",
                    "land_number": "",
                    "owner_name": "",
                    "area": "abc",
                    "denominator": "0",
                }
            ],
            enabled_rules=["land_number_missing", "denominator_zero"],
        )

        categories = {issue.category for issue in issues}
        self.assertEqual(categories, {"地號缺漏", "分母為 0"})

    def test_detects_probable_duplicate_same_land_and_owner(self):
        records = [
            {
                "id": 1,
                "district": "桃園市",
                "section": "一段",
                "registration_order": "15",
                "land_number": "100",
                "owner_name": "王弘益",
                "external_id": "H100059743",
            },
            {
                "id": 2,
                "district": "桃園市",
                "section": "一段",
                "registration_order": "15",
                "land_number": "100",
                "owner_name": "王弘益",
                "external_id": "H100059743",
            },
            {
                "id": 3,
                "district": "桃園市",
                "section": "一段",
                "registration_order": "16",
                "land_number": "100",
                "owner_name": "不同所有權人",
                "external_id": "",
            },
        ]

        issues = inspect_customer_quality(records)

        duplicate_ids = {
            issue.record_id
            for issue in issues
            if issue.category == "疑似重複"
        }
        self.assertEqual(duplicate_ids, {1, 2})

    def test_quality_issue_signature_is_stable_and_specific(self):
        issue = DataQualityIssue(
            record_id=7,
            severity="重要",
            category="分母為 0",
            summary="#7 桃園市 / 一段 / 100：分母不可為 0",
            suggestion="請修正分母。",
        )
        same_issue = DataQualityIssue(
            record_id=7,
            severity="重要",
            category="分母為 0",
            summary="#7 桃園市 / 一段 / 100：分母不可為 0",
            suggestion="請修正分母。",
        )
        changed_issue = DataQualityIssue(
            record_id=7,
            severity="重要",
            category="分母為 0",
            summary="#7 桃園市 / 一段 / 101：分母不可為 0",
            suggestion="請修正分母。",
        )

        self.assertEqual(quality_issue_signature(issue), quality_issue_signature(same_issue))
        self.assertNotEqual(quality_issue_signature(issue), quality_issue_signature(changed_issue))

    def test_write_quality_issues_creates_readable_excel(self):
        with tempfile.TemporaryDirectory() as temp_name:
            file_path = Path(temp_name) / "quality.xlsx"
            write_quality_issues(
                file_path,
                [
                    DataQualityIssue(
                        record_id=7,
                        severity="重要",
                        category="分母為 0",
                        summary="#7 桃園市 / 一段 / 100：分母不可為 0",
                        suggestion="請修正分母。",
                    )
                ],
            )

            workbook = load_workbook(file_path)
            try:
                sheet = workbook.active
                self.assertEqual(sheet.title, "資料品質檢查")
                self.assertEqual(
                    [sheet.cell(row=1, column=column).value for column in range(1, 6)],
                    ["資料ID", "程度", "類型", "問題", "建議處理"],
                )
                self.assertEqual(sheet.cell(row=2, column=1).value, 7)
                self.assertEqual(sheet.cell(row=2, column=3).value, "分母為 0")
            finally:
                workbook.close()

    def test_write_quality_summary_creates_readable_excel(self):
        with tempfile.TemporaryDirectory() as temp_name:
            file_path = Path(temp_name) / "quality-summary.xlsx"
            write_quality_summary(
                file_path,
                [
                    {"category": "分母為 0", "important": 1, "reminder": 0, "total": 1},
                    {"category": "已確認隱藏", "important": "", "reminder": "", "total": 2},
                ],
            )

            workbook = load_workbook(file_path)
            try:
                sheet = workbook.active
                self.assertEqual(sheet.title, "品質檢查摘要")
                self.assertEqual(
                    [sheet.cell(row=1, column=column).value for column in range(1, 5)],
                    ["類型", "重要", "提醒", "合計"],
                )
                self.assertEqual(sheet.cell(row=2, column=1).value, "分母為 0")
                self.assertEqual(sheet.cell(row=2, column=2).value, 1)
                self.assertEqual(sheet.cell(row=3, column=1).value, "已確認隱藏")
                self.assertEqual(sheet.cell(row=3, column=4).value, 2)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
