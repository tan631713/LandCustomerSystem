import os
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook, load_workbook
from PySide6.QtCore import QEventLoop, QThread, QTimer
from PySide6.QtWidgets import QApplication

import customer_ui_qt as app
from customer_excel import ExcelExportWorker, ExcelImportWorker, ExcelService


class CustomerExcelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.application = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp_context = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_context.name)
        self.service = ExcelService(
            app.LAND_FIELDS,
            app.parse_number,
            app.split_rights_scope,
            app.format_number_text,
            app.calculate_ping,
            app.format_ping_text,
            app.calculate_total_declared_value,
        )

    def tearDown(self):
        self.temp_context.cleanup()

    def create_import_file(self):
        path = self.root / "import.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["地區", "地段", "地號", "面積/m2", "公告現值", "分子", "分母", "姓名"])
        sheet.append(["中正區", "一段", "100", 100, 20000, 1, 2, "王小明"])
        sheet.append([None, "一段", "101", 80, 18000, 1, 0, "陳小華"])
        workbook.save(path)
        workbook.close()
        return path

    def test_import_recognizes_columns_forward_fills_and_reports_invalid_rows(self):
        result = self.service.read_import_file(self.create_import_file())

        self.assertEqual(len(result["records"]), 1)
        self.assertEqual(len(result["error_rows"]), 1)
        self.assertEqual(result["records"][0]["district"], "中正區")
        self.assertEqual(result["records"][0]["total_declared_value"], "1,000,000")
        self.assertIn("分母不能為 0", result["error_rows"][0]["reason"])
        self.assertEqual(result["error_rows"][0]["data"]["district"], "中正區")

    def test_export_writes_requested_columns_and_values(self):
        path = self.root / "export.xlsx"
        rows = [{"district": "中正區", "owner_name": "王小明"}]
        columns = [("district", "地區"), ("owner_name", "姓名")]

        self.service.write_rows(path, rows, columns)

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            self.assertEqual(sheet.cell(1, 1).value, "地區")
            self.assertEqual(sheet.cell(2, 2).value, "王小明")
        finally:
            workbook.close()

    def test_workers_emit_success_results(self):
        import_results = []
        import_errors = []
        import_worker = ExcelImportWorker(self.service, self.create_import_file())
        import_worker.finished.connect(import_results.append)
        import_worker.failed.connect(import_errors.append)
        import_worker.run()

        export_results = []
        export_errors = []
        export_path = self.root / "worker-export.xlsx"
        export_worker = ExcelExportWorker(
            export_path,
            [{"district": "信義區"}],
            [("district", "地區")],
        )
        export_worker.finished.connect(lambda path, count: export_results.append((path, count)))
        export_worker.failed.connect(export_errors.append)
        export_worker.run()

        self.assertFalse(import_errors)
        self.assertEqual(len(import_results[0]["records"]), 1)
        self.assertFalse(export_errors)
        self.assertEqual(export_results[0][1], 1)
        self.assertTrue(export_path.exists())

    def test_export_worker_runs_on_qthread(self):
        export_path = self.root / "thread-export.xlsx"
        worker = ExcelExportWorker(
            export_path,
            [{"district": "大安區"}],
            [("district", "地區")],
        )
        thread = QThread()
        loop = QEventLoop()
        results, errors = [], []
        worker.moveToThread(thread)
        thread.started.connect(worker.run)
        worker.finished.connect(lambda path, count: results.append((path, count)))
        worker.finished.connect(thread.quit)
        worker.failed.connect(errors.append)
        worker.failed.connect(thread.quit)
        thread.finished.connect(loop.quit)
        QTimer.singleShot(5000, loop.quit)

        thread.start()
        loop.exec()
        thread.wait(1000)

        self.assertFalse(thread.isRunning())
        self.assertFalse(errors)
        self.assertEqual(results[0][1], 1)
        self.assertTrue(export_path.exists())

    def test_excel_round_trip_without_optional_numpy_or_lxml(self):
        project_root = Path(__file__).resolve().parents[1]
        script = r'''
import importlib.abc
import tempfile
from pathlib import Path

class OptionalDependencyBlocker(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname.split(".", 1)[0] in {"numpy", "lxml"}:
            raise ModuleNotFoundError(fullname)
        return None

import sys
sys.meta_path.insert(0, OptionalDependencyBlocker())
from customer_excel import ExcelService

with tempfile.TemporaryDirectory() as temp_dir:
    output = Path(temp_dir) / "round-trip.xlsx"
    ExcelService.write_rows(output, [{"district": "D"}], [("district", "District")])
    from openpyxl import load_workbook
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        assert workbook.active.cell(2, 1).value == "D"
    finally:
        workbook.close()
'''
        result = subprocess.run(
            [sys.executable, "-c", script],
            cwd=project_root,
            capture_output=True,
            text=True,
            timeout=20,
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_excel_round_trip_with_packaged_numpy_missing_short_alias(self):
        project_root = Path(__file__).resolve().parents[1]
        script = r'''
import sys
import tempfile
import types
from pathlib import Path

numpy = types.ModuleType("numpy")
sys.modules["numpy"] = numpy

from customer_excel import ExcelService

with tempfile.TemporaryDirectory() as temp_dir:
    output = Path(temp_dir) / "round-trip.xlsx"
    ExcelService.write_rows(output, [{"district": "D"}], [("district", "District")])
    from openpyxl import load_workbook
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        assert workbook.active.cell(2, 1).value == "D"
        assert hasattr(numpy, "short")
    finally:
        workbook.close()
'''
        result = subprocess.run(
            [sys.executable, "-c", script],
            cwd=project_root,
            capture_output=True,
            text=True,
            timeout=20,
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_large_excel_import_and_export_complete_within_limit(self):
        import_path = self.root / "large-import.xlsx"
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet()
        sheet.append(["地區", "地段", "地號", "面積/m2", "公告現值", "分子", "分母", "姓名"])
        for number in range(2000):
            sheet.append(["中正區", "一段", str(number), 100, 20000, 1, 2, f"姓名{number}"])
        workbook.save(import_path)
        workbook.close()

        started = time.perf_counter()
        result = self.service.read_import_file(import_path)
        import_elapsed = time.perf_counter() - started
        self.assertEqual(len(result["records"]), 2000)
        self.assertFalse(result["error_rows"])
        self.assertLess(import_elapsed, 15.0)

        export_path = self.root / "large-export.xlsx"
        export_rows = [
            {"district": "中正區", "land_number": str(number)}
            for number in range(5000)
        ]
        started = time.perf_counter()
        self.service.write_rows(
            export_path,
            export_rows,
            [("district", "地區"), ("land_number", "地號")],
        )
        export_elapsed = time.perf_counter() - started
        self.assertLess(export_elapsed, 15.0)
        workbook = load_workbook(export_path, read_only=True, data_only=True)
        try:
            self.assertEqual(workbook.active.max_row, 5001)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
