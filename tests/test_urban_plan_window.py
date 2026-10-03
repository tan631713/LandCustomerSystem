"""The 都市計畫 view, 輸入到 selector and plan dialogs inside the real window."""

import os
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QEventLoop, QTimer, Qt
from PySide6.QtWidgets import QApplication

import customer_import_controller as import_controller
import customer_ui_qt as app
import customer_urban_plan_dialogs as plan_dialogs
from customer_database import CustomerDatabase
from customer_desktop_api import DesktopApiRecordRepository, DesktopApiResponseError
from customer_repository import CustomerRepository
from customer_security import make_fernet
from customer_urban_plan_dialogs import SetUrbanPlanDialog, UrbanPlanManagementDialog

LONGGANG = (1, "龍岡都市計畫")
DANAN = (2, "大湳都市計畫")


def make_rows():
    """Three lands in 龍岡, one in 大湳 and two unclassified (nine ownerships)."""

    layout = [
        # land_id, plan, section, land_number, owners
        (10, LONGGANG, "龍岡段", "221-0000", ["王一", "王二"]),
        (11, LONGGANG, "龍岡段", "305-0000", ["李三"]),
        (12, LONGGANG, "忠福段", "9-0000", ["陳四", "陳五"]),
        (20, DANAN, "大湳段", "7-0000", ["張六"]),
        (30, None, "未定段", "1-0000", ["黃七", "黃八"]),
        (31, None, "未定段", "2-0000", ["林九"]),
    ]
    rows = []
    record_id = 0
    for land_id, plan, section, land_number, owners in layout:
        for owner in owners:
            record_id += 1
            rows.append({
                "id": record_id, "ownership_id": record_id, "land_id": land_id, "owner_id": record_id,
                "district": "中壢區", "section": section, "subsection": "",
                "registration_order": f"{record_id:04d}", "land_number": land_number,
                "area": "100", "declared_value": "1000", "numerator": "1", "denominator": "2",
                "ping": "30.25", "total_declared_value": "500", "registration_reason": "買賣",
                "note": "", "visit_log": "", "owner_name": owner, "name": owner,
                "external_id": f"A{100000000 + record_id}", "address": "某路 1 號",
                "birth_year": "1960", "created_at": "2026-01-01T00:00:00",
                "updated_at": "2026-01-01T00:00:00", "case_names": "", "tag_names": "",
                "tag_items": [], "tags": [], "primary_tag_color": "", "attachment_count": 0,
                "attachment_names": "", "custom_values": "",
                "urban_plan_id": plan[0] if plan else None,
                "urban_plan_name": plan[1] if plan else None,
            })
    rows.sort(key=lambda row: row["id"], reverse=True)
    return rows


class PlanHomeServer:
    """A fake home server with working 都市計畫 endpoints."""

    def __init__(self, *, supports_plans=True):
        self.rows = make_rows()
        self.supports_plans = supports_plans
        self.plans = {1: "龍岡都市計畫", 2: "大湳都市計畫", 3: "空的都市計畫"}
        self.saved = []
        self.assignments = []
        self.imports = []
        self.requests = []

    def _hit(self, name):
        self.requests.append((name, threading.current_thread() is threading.main_thread()))

    # -- records ---------------------------------------------------------------
    def list_all_records(self, **_kwargs):
        return [dict(row) for row in self.rows]

    def get_record(self, record_id):
        return dict(next(row for row in self.rows if row["id"] == int(record_id)))

    def _store(self, record_id, values):
        values = dict(values)
        plan_id = values.pop("urban_plan_id", None)
        for row in self.rows:
            if row["id"] == int(record_id):
                row.update(values)
                land_id = row["land_id"]
        if plan_id is not None:
            self._set_land_plan(land_id, plan_id or None)

    def _set_land_plan(self, land_id, plan_id):
        for row in self.rows:
            if row["land_id"] == land_id:
                row["urban_plan_id"] = plan_id
                row["urban_plan_name"] = self.plans.get(plan_id) if plan_id else None

    def replace_record_with_history(self, record_id, values, _logs):
        self._hit("save record")
        self.saved.append(dict(values))
        self._store(record_id, values)
        return int(record_id)

    def replace_record(self, record_id, values):
        return self.replace_record_with_history(record_id, values, [])

    def create_record(self, values):
        self._hit("create record")
        self.saved.append(dict(values))
        record_id = max(row["id"] for row in self.rows) + 1
        row = dict(self.rows[0])
        plan_id = values.get("urban_plan_id")
        row.update({key: value for key, value in values.items() if key != "urban_plan_id"})
        row.update(id=record_id, ownership_id=record_id, land_id=900 + record_id, owner_id=record_id,
                   urban_plan_id=plan_id or None,
                   urban_plan_name=self.plans.get(plan_id) if plan_id else None)
        self.rows.insert(0, row)
        return record_id

    # -- plans --------------------------------------------------------------------
    def _counts(self, plan_id):
        rows = [row for row in self.rows if (row["urban_plan_id"] or 0) == plan_id]
        return {
            "land_count": len({row["land_id"] for row in rows}),
            "ownership_count": len(rows),
            "owner_count": len({row["owner_id"] for row in rows}),
        }

    def list_urban_plans(self):
        self._hit("list plans")
        if not self.supports_plans:
            raise DesktopApiResponseError(404, "找不到 API 路徑")
        return {
            "items": [
                {"id": plan_id, "plan_id": plan_id, "name": name, **self._counts(plan_id)}
                for plan_id, name in sorted(self.plans.items())
            ],
            "unassigned": self._counts(0),
        }

    def save_urban_plan(self, name, plan_id=None):
        if any(existing.casefold() == name.casefold() and key != plan_id
               for key, existing in self.plans.items()):
            raise DesktopApiResponseError(409, "已有相同名稱的都市計畫")
        plan_id = plan_id or max(self.plans, default=0) + 1
        self.plans[plan_id] = name
        for row in self.rows:
            if row["urban_plan_id"] == plan_id:
                row["urban_plan_name"] = name
        return plan_id

    def delete_urban_plan(self, plan_id):
        released = {row["land_id"] for row in self.rows if row["urban_plan_id"] == plan_id}
        for land_id in released:
            self._set_land_plan(land_id, None)
        self.plans.pop(plan_id, None)
        return {"deleted": True, "released_land_count": len(released)}

    def set_lands_urban_plan(self, land_ids, plan_id, only_unassigned=True):
        self.assignments.append((sorted(land_ids), plan_id, only_unassigned))
        updated = []
        for land_id in sorted(set(land_ids)):
            current = next(row["urban_plan_id"] for row in self.rows if row["land_id"] == land_id)
            if plan_id and only_unassigned and current:
                continue
            self._set_land_plan(land_id, plan_id or None)
            updated.append(land_id)
        return {"updated_land_ids": updated, "updated_count": len(updated), "kept_lands": []}

    def import_records(self, items, source_file_name, urban_plan_id=None):
        """Insert each item (a land with the same district/section/number is the
        same land); only lands without a plan are put in `urban_plan_id`."""

        self.imports.append((len(items), urban_plan_id))
        inserted, touched = [], []
        for item in items:
            values = dict(item["values"])
            record_id = max(row["id"] for row in self.rows) + 1
            same_land = next(
                (row for row in self.rows
                 if (row["district"], row["section"], row["land_number"])
                 == (values.get("district"), values.get("section"), values.get("land_number"))),
                None,
            )
            row = dict(self.rows[0])
            row.update(values)
            row.update(
                id=record_id, ownership_id=record_id, owner_id=record_id,
                land_id=same_land["land_id"] if same_land else 1000 + record_id,
                urban_plan_id=same_land["urban_plan_id"] if same_land else None,
                urban_plan_name=same_land["urban_plan_name"] if same_land else None,
            )
            self.rows.insert(0, row)
            inserted.append(record_id)
            if row["land_id"] not in touched:
                touched.append(row["land_id"])
        result = {"batch_id": 1, "inserted_ids": inserted, "updated_ids": [],
                  "inserted_count": len(inserted), "updated_count": 0}
        if urban_plan_id:
            assigned, kept = [], []
            for land_id in touched:
                land_row = next(row for row in self.rows if row["land_id"] == land_id)
                if land_row["urban_plan_id"]:
                    kept.append({"id": land_id, "district": land_row["district"],
                                 "section": land_row["section"], "subsection": "",
                                 "land_number": land_row["land_number"],
                                 "urban_plan_name": land_row["urban_plan_name"]})
                else:
                    self._set_land_plan(land_id, urban_plan_id)
                    assigned.append(land_id)
            result["urban_plan"] = {
                "id": urban_plan_id, "name": self.plans[urban_plan_id],
                "assigned_land_count": len(assigned), "kept_land_count": len(kept),
                "kept_lands": kept,
            }
        return result

    def list_watchlist(self):
        return []

    def add_operation_log(self, *_args, **_kwargs):
        return 1

    def __getattr__(self, name):
        def any_other_request(*_args, **_kwargs):
            self._hit(name)
            return []

        return any_other_request


class UrbanPlanWindowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.application = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp_context = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_context.name)
        self.originals = {
            "DB_PATH": app.DB_PATH, "BACKUP_DIR": app.BACKUP_DIR,
            "DATABASE": app.DATABASE, "REPOSITORY": app.REPOSITORY,
        }
        app.DB_PATH = self.root / "customers.db"
        app.BACKUP_DIR = self.root / "backups"
        app.DATABASE = CustomerDatabase(app.DB_PATH, app.BACKUP_DIR)
        app.REPOSITORY = CustomerRepository(
            app.DATABASE, app.SCHEMA_PATH, self.root / "missing-seed.sql", app.LAND_FIELDS
        )
        app.init_db()
        app.REPOSITORY.create_admin_user("test-password")
        self.key = app.REPOSITORY.authenticate_user("admin", "test-password")
        self.windows = []
        # A modal box nobody answers would hang the whole run; tests that care
        # about a message patch it again and look at the mock.
        for module in (app, plan_dialogs, import_controller):
            patcher = patch.object(module, "QMessageBox")
            patcher.start()
            self.addCleanup(patcher.stop)

    def tearDown(self):
        for window in self.windows:
            window.close()
        self.application.processEvents()
        for name, value in self.originals.items():
            setattr(app, name, value)
        self.temp_context.cleanup()

    def open_window(self, *, supports_plans=True, settings=None):
        for key, value in (settings or {}).items():
            app.REPOSITORY.set_setting(key, value)
        server = PlanHomeServer(supports_plans=supports_plans)
        repository = DesktopApiRecordRepository(server, make_fernet(self.key))
        window = app.LandApp(self.key, record_repository=repository, api_mode=True)
        self.windows.append(window)
        self.settle(window)
        return server, repository, window

    def settle(self, window, timeout_ms=15000):
        deadline = QTimer()
        deadline.setSingleShot(True)
        loop = QEventLoop()
        deadline.timeout.connect(loop.quit)
        deadline.start(timeout_ms)
        poll = QTimer()

        def check():
            window.wait_for_background_tasks(0.01)
            if not (window.record_searches or getattr(window, "_saved_record_reconcile", None)):
                loop.quit()

        poll.timeout.connect(check)
        poll.start(20)
        loop.exec()
        poll.stop()
        deadline.stop()
        window.wait_for_background_tasks()
        self.application.processEvents()

    def show_plan_view(self, window, mode="all"):
        index = window.plan_view_combo.findData(mode)
        self.assertGreaterEqual(index, 0, mode)
        window.plan_view_combo.setCurrentIndex(index)
        self.settle(window)

    @staticmethod
    def titles(window):
        view = window.urban_plan_view
        return [view.topLevelItem(i).node.title for i in range(view.topLevelItemCount())]

    # -- controls -------------------------------------------------------------------------

    def test_controls_appear_with_a_server_that_has_plans(self):
        server, _repository, window = self.open_window()
        self.assertTrue(window.urban_plans_available())
        self.assertFalse(window.plan_view_combo.isHidden())
        self.assertFalse(window.plan_input_combo.isHidden())
        self.assertFalse(window.form_field_containers["urban_plan"].isHidden())
        self.assertEqual(
            [window.plan_view_combo.itemText(i) for i in range(window.plan_view_combo.count())],
            ["依土地（舊）", "全部都市計畫", "龍岡都市計畫", "大湳都市計畫", "空的都市計畫", "未分類"],
        )
        self.assertEqual(
            [window.plan_input_combo.itemText(i) for i in range(window.plan_input_combo.count())],
            ["不指定（未分類）", "龍岡都市計畫", "大湳都市計畫", "空的都市計畫"],
        )
        self.assertEqual(window.table_stack.currentIndex(), 0)  # the land tree stays the default
        self.assertEqual(window.expand_all_button.text(), "全部展開")

    def test_an_older_server_without_plans_hides_everything_and_saves_normally(self):
        server, _repository, window = self.open_window(supports_plans=False)
        self.assertFalse(window.urban_plans_available())
        for widget in (window.plan_view_combo, window.plan_input_combo):
            self.assertTrue(widget.isHidden())
        self.assertTrue(window.form_field_containers["urban_plan"].isHidden())
        window.load_record(3)
        window.set_field_text("note", "舊伺服器")
        with patch.object(app, "QMessageBox"):
            window.save_record()
        self.settle(window)
        self.assertEqual(len(server.saved), 1)
        self.assertNotIn("urban_plan_id", server.saved[0])

    def test_standalone_mode_has_no_plan_controls(self):
        window = app.LandApp(self.key)
        self.windows.append(window)
        self.assertFalse(window.urban_plans_available())
        self.assertTrue(window.plan_view_combo.isHidden())
        self.assertEqual(window.table_stack.currentIndex(), 0)

    # -- the plan view ------------------------------------------------------------------------

    def test_all_plans_view_groups_lands_and_hides_the_page_controls(self):
        server, _repository, window = self.open_window()
        self.show_plan_view(window)
        self.assertEqual(window.table_stack.currentIndex(), 1)
        self.assertEqual(self.titles(window), ["大湳都市計畫", "空的都市計畫", "龍岡都市計畫", "未分類"])
        self.assertTrue(window.next_land_page_button.isHidden())
        self.assertTrue(window.land_page_label.isHidden())
        tree = window.urban_plan_tree
        self.assertEqual((tree.plan_count, tree.land_count, tree.share_count, tree.owner_count), (3, 6, 9, 9))
        self.assertIn("都市計畫 3 個", window.land_count_label.text())
        self.assertIn("土地 6 筆", window.land_count_label.text())
        longgang = window.urban_plan_view.topLevelItem(2)
        self.assertTrue(longgang.isExpanded())  # plans open by default
        self.assertFalse(longgang.child(0).isExpanded())  # sections stay closed
        self.assertEqual([longgang.text(i) for i in (1, 2)], ["5", "5"])

    def test_choosing_one_plan_shows_only_that_plan(self):
        server, _repository, window = self.open_window()
        self.show_plan_view(window, "plan:1")
        self.assertEqual(self.titles(window), ["龍岡都市計畫"])
        self.assertEqual(window.urban_plan_tree.land_count, 3)
        self.assertIn("龍岡都市計畫", window.land_count_label.text())
        self.assertIn("只顯示「龍岡都市計畫」", window.plan_hint_label.text())
        section = window.urban_plan_view.topLevelItem(0).child(0)
        self.assertTrue(section.isExpanded())  # a single plan opens one level deeper
        self.show_plan_view(window, "plan:0")
        self.assertEqual(self.titles(window), ["未分類"])
        self.assertEqual(window.urban_plan_tree.land_count, 2)
        self.show_plan_view(window, "plan:3")
        self.assertEqual(self.titles(window), ["空的都市計畫"])
        self.show_plan_view(window, "legacy")
        self.assertEqual(window.table_stack.currentIndex(), 0)
        self.assertFalse(window.next_land_page_button.isHidden())

    def test_the_chosen_view_and_input_plan_are_remembered(self):
        _server, _repository, window = self.open_window()
        self.show_plan_view(window, "plan:2")
        window.plan_input_combo.setCurrentIndex(window.plan_input_combo.findData(1))
        window.close()
        _server, _repository, again = self.open_window()
        self.assertEqual(again.urban_plan_mode, ("plan", 2))
        self.assertEqual(again.table_stack.currentIndex(), 1)
        self.assertEqual(again.urban_plan_input_id, 1)
        self.assertEqual(again.plan_input_combo.currentData(), 1)

    def test_a_remembered_plan_that_no_longer_exists_falls_back_to_all(self):
        _server, _repository, window = self.open_window(settings={"urban_plan_view": "plan:99"})
        self.assertEqual(window.urban_plan_mode, ("all", None))

    def test_clicking_an_owner_loads_the_form_with_the_plan(self):
        _server, _repository, window = self.open_window()
        self.show_plan_view(window)
        node = window.urban_plan_tree.node_for_record(3)  # 李三, in 龍岡
        window.on_plan_node_selected(node)
        self.settle(window)
        self.assertEqual(window.selected_record_id, 3)
        self.assertEqual(window.get_field_text("owner_name"), "李三")
        self.assertEqual(window.urban_plan_form_value(), 1)
        self.assertEqual(window.urban_plan_form_combo.currentText(), "龍岡都市計畫")
        unclassified = window.urban_plan_tree.node_for_record(9)
        window.on_plan_node_selected(unclassified)
        self.settle(window)
        self.assertEqual(window.urban_plan_form_value(), 0)

    # -- checkboxes -------------------------------------------------------------------------------

    def test_checking_a_plan_checks_every_ownership_below_it_and_back(self):
        _server, _repository, window = self.open_window()
        self.show_plan_view(window)
        view = window.urban_plan_view
        longgang = view.topLevelItem(2)
        longgang.setCheckState(0, Qt.Checked)  # what a click on the checkbox does
        self.application.processEvents()
        self.assertEqual(sorted(window.checked_record_ids), [1, 2, 3, 4, 5])
        self.assertEqual(longgang.checkState(0), Qt.Checked)
        self.assertEqual(longgang.child(0).checkState(0), Qt.Checked)
        window.checked_record_ids.discard(2)
        window.urban_plan_checks_changed()
        self.assertEqual(longgang.checkState(0), Qt.PartiallyChecked)
        self.assertEqual(view.topLevelItem(0).checkState(0), Qt.Unchecked)
        longgang.setCheckState(0, Qt.Unchecked)
        self.assertEqual(sorted(window.checked_record_ids), [])

    def test_legacy_check_actions_update_the_plan_view(self):
        _server, _repository, window = self.open_window()
        self.show_plan_view(window)
        window.update_checked_records([6], True, "測試")
        owner_item = window.urban_plan_view._record_items[6]
        self.assertEqual(owner_item.checkState(0), Qt.Checked)
        window.clear_checked_selection()
        self.assertEqual(owner_item.checkState(0), Qt.Unchecked)

    def test_selection_helpers_follow_the_plan_view(self):
        _server, _repository, window = self.open_window()
        self.show_plan_view(window)
        view = window.urban_plan_view
        view.topLevelItem(0).setSelected(True)  # 大湳都市計畫
        self.assertEqual(window.selected_table_record_ids(), [6])
        section = view.topLevelItem(2).child(1)  # 龍岡 / 龍岡段: lands 221 and 305
        view.clearSelection()
        section.setSelected(True)
        self.assertEqual(window.selected_table_record_ids(), [1, 2, 3])

    # -- search ----------------------------------------------------------------------------------

    def test_searching_opens_everything_and_hides_plans_without_matches(self):
        _server, _repository, window = self.open_window()
        self.show_plan_view(window)
        window.search_input.setText("張六")
        window.refresh_records_for_search()
        self.settle(window)
        self.assertEqual(self.titles(window), ["大湳都市計畫"])
        owner_item = window.urban_plan_view._record_items[6]
        self.assertTrue(owner_item.parent().isExpanded())
        self.assertTrue(owner_item.parent().parent().isExpanded())
        window.search_input.setText("")
        window.refresh_records_for_search()
        self.settle(window)
        self.assertEqual(len(self.titles(window)), 4)

    # -- saving --------------------------------------------------------------------------------------

    def save(self, window, server):
        server.requests.clear()
        with patch.object(app, "QMessageBox"):
            window.save_record()
        self.settle(window)

    def test_changing_the_plan_in_the_form_moves_the_whole_land(self):
        server, _repository, window = self.open_window()
        self.show_plan_view(window)
        window.load_record(7)  # 黃七, unclassified land 30 (two owners)
        self.settle(window)
        window.urban_plan_form_combo.setCurrentIndex(window.urban_plan_form_combo.findData(2))
        self.save(window, server)
        self.assertEqual(server.saved[-1]["urban_plan_id"], 2)
        self.assertEqual([row["urban_plan_id"] for row in server.rows if row["land_id"] == 30], [2, 2])
        self.assertEqual(server.requests[0], ("save record", True))
        tree = window.urban_plan_tree
        danan = next(root for root in tree.roots if root.plan_id == 2)
        self.assertEqual(danan.share_count, 3)
        self.assertIn(window.selected_record_id, danan.record_ids)

    def test_saving_without_touching_the_plan_does_not_send_it(self):
        server, _repository, window = self.open_window()
        window.load_record(4)
        self.settle(window)
        window.set_field_text("note", "只改備註")
        self.save(window, server)
        self.assertNotIn("urban_plan_id", server.saved[-1])

    def test_saving_never_waits_for_the_plan_list(self):
        server, _repository, window = self.open_window()
        window.load_record(4)
        self.settle(window)
        window.set_field_text("note", "x")
        self.save(window, server)
        self.assertNotIn("list plans", [name for name, _on_gui in server.requests])

    def test_a_new_record_starts_in_the_input_plan_and_is_saved_into_it(self):
        server, _repository, window = self.open_window()
        window.plan_input_combo.setCurrentIndex(window.plan_input_combo.findData(2))
        window.new_record()
        self.assertEqual(window.urban_plan_form_value(), 2)
        for key, value in {"district": "中壢區", "section": "新段", "land_number": "77-0000",
                           "owner_name": "新地主"}.items():
            window.set_field_text(key, value)
        self.save(window, server)
        self.assertEqual(server.saved[-1]["urban_plan_id"], 2)
        new_row = server.rows[0]
        self.assertEqual(new_row["urban_plan_name"], "大湳都市計畫")

    def test_changing_input_plan_while_editing_an_existing_record_keeps_its_own_plan(self):
        _server, _repository, window = self.open_window()
        window.load_record(4)
        window.plan_input_combo.setCurrentIndex(window.plan_input_combo.findData(2))
        self.assertEqual(window.urban_plan_form_value(), 1)

    # -- dialogs -----------------------------------------------------------------------------------------

    def test_management_dialog_creates_renames_and_deletes_plans(self):
        server, repository, window = self.open_window()
        window.manage_urban_plans()
        dialog = window.urban_plan_management_dialog
        self.assertIsInstance(dialog, UrbanPlanManagementDialog)
        names = [dialog.table.item(row, 0).text() for row in range(dialog.table.rowCount())]
        self.assertEqual(names, ["龍岡都市計畫", "大湳都市計畫", "空的都市計畫", "未分類"])
        self.assertEqual(dialog.table.item(0, 1).text(), "3")  # lands in 龍岡
        dialog.name_input.setText("  新  計畫 ")
        self.assertTrue(dialog.add_plan())
        self.assertEqual(server.plans[4], "新 計畫")
        with patch.object(plan_dialogs, "QMessageBox") as box:
            dialog.name_input.setText("龍岡都市計畫")
            self.assertFalse(dialog.add_plan())
            self.assertIn("已有相同名稱", box.warning.call_args[0][2])
        dialog.table.selectRow(0)
        with patch.object(plan_dialogs.QInputDialog, "getText", return_value=("龍岡改名", True)):
            self.assertTrue(dialog.rename_selected())
        self.assertEqual(server.plans[1], "龍岡改名")
        dialog.table.selectRow(dialog.table.rowCount() - 1)  # 未分類 cannot be edited
        self.assertFalse(dialog.rename_button.isEnabled())
        self.assertFalse(dialog.delete_button.isEnabled())
        dialog.close()
        self.settle(window)
        self.assertIn("龍岡改名", [window.plan_view_combo.itemText(i) for i in range(window.plan_view_combo.count())])
        self.assertIn("新 計畫", [window.plan_input_combo.itemText(i) for i in range(window.plan_input_combo.count())])
        self.assertIn("龍岡改名", [window.urban_plan_form_combo.itemText(i) for i in range(window.urban_plan_form_combo.count())])

    def test_deleting_a_plan_releases_its_lands_and_resets_selectors(self):
        server, repository, window = self.open_window()
        self.show_plan_view(window, "plan:1")
        window.plan_input_combo.setCurrentIndex(window.plan_input_combo.findData(1))
        window.manage_urban_plans()
        dialog = window.urban_plan_management_dialog
        dialog.table.selectRow(0)
        with patch.object(plan_dialogs.DeleteUrbanPlanDialog, "exec", return_value=plan_dialogs.QDialog.Accepted):
            self.assertTrue(dialog.delete_selected())
        self.assertNotIn(1, server.plans)
        self.assertEqual([row["urban_plan_id"] for row in server.rows if row["land_id"] == 10], [None, None])
        dialog.close()
        self.settle(window)
        self.assertEqual(window.urban_plan_mode, ("all", None))  # the viewed plan is gone
        self.assertEqual(window.urban_plan_input_id, 0)
        self.assertEqual(window.urban_plan_tree.land_count, 6)
        self.assertEqual(
            next(root for root in window.urban_plan_tree.roots if root.plan_id == 0).share_count, 8
        )

    def test_creating_the_first_plan_switches_from_the_old_view_to_the_plan_view(self):
        server, repository, window = self.open_window()
        server.plans.clear()
        for row in server.rows:
            row["urban_plan_id"] = row["urban_plan_name"] = None
        repository.list_urban_plans(refresh=True)
        window.reload_urban_plans()
        window.manage_urban_plans()
        dialog = window.urban_plan_management_dialog
        dialog.name_input.setText("龍岡都市計畫")
        dialog.add_plan()
        dialog.close()
        self.settle(window)
        self.assertEqual(window.urban_plan_mode, ("all", None))
        self.assertEqual(window.table_stack.currentIndex(), 1)

    def test_batch_assignment_leaves_other_plans_alone_by_default(self):
        server, repository, window = self.open_window()
        window.new_record()  # nothing selected: only the checked rows count
        # Ownership 1 is on land 10 (龍岡); ownership 7 is on land 30 (未分類).
        window.checked_record_ids.update({1, 7})
        with patch.object(app, "QMessageBox"):
            window.set_urban_plan_for_selection()
            dialog = next(d for d in window._open_edit_dialogs if isinstance(d, SetUrbanPlanDialog))
            self.assertTrue(dialog.only_unassigned_check.isChecked())
            self.assertEqual(dialog.land_ids(), [10, 30])
            dialog.plan_combo.setCurrentIndex(dialog.plan_combo.findData(2))
            dialog.accept()
            self.settle(window)
        self.assertEqual(server.assignments[-1], ([10, 30], 2, True))
        plans_by_land = {row["land_id"]: row["urban_plan_id"] for row in server.rows}
        self.assertEqual(plans_by_land[10], 1)  # already in 龍岡: untouched
        self.assertEqual(plans_by_land[30], 2)

    def test_batch_assignment_needs_a_selection_and_at_least_one_plan(self):
        server, repository, window = self.open_window()
        window.new_record()
        with patch.object(app, "QMessageBox") as box:
            window.set_urban_plan_for_selection()
            self.assertTrue(box.warning.called)
        self.assertEqual(server.assignments, [])

    def test_batch_assignment_dialog_counts(self):
        plans = [{"plan_id": 1, "name": "龍岡都市計畫"}, {"plan_id": 2, "name": "大湳都市計畫"}]
        dialog = SetUrbanPlanDialog(plans, {10: 1, 11: 0, 12: 2, 13: 0}, 6)
        dialog.plan_combo.setCurrentIndex(dialog.plan_combo.findData(1))
        self.assertEqual(dialog.affected_land_count(), 2)  # only the two unclassified
        self.assertIn("將有 2 筆土地設定為「龍岡都市計畫」", dialog.summary_label.text())
        self.assertIn("1 筆已屬於其他都市計畫，不會變更", dialog.summary_label.text())
        dialog.only_unassigned_check.setChecked(False)
        self.assertEqual(dialog.affected_land_count(), 3)
        self.assertFalse(dialog.only_unassigned())
        self.assertIn("1 筆原屬其他都市計畫，會被改到這個計畫", dialog.summary_label.text())
        dialog.plan_combo.setCurrentIndex(dialog.plan_combo.findData(0))
        self.assertEqual(dialog.affected_land_count(), 2)  # releasing: the two classified lands
        self.assertFalse(dialog.only_unassigned_check.isEnabled())
        self.assertEqual(dialog.land_ids(), [10, 11, 12, 13])
        nothing = SetUrbanPlanDialog(plans, {10: 1}, 1)
        nothing.plan_combo.setCurrentIndex(nothing.plan_combo.findData(1))
        self.assertFalse(nothing.apply_button.isEnabled())

    # -- importing ------------------------------------------------------------------------------------

    @staticmethod
    def import_result():
        def record(section, land_number, owner):
            return {"district": "中壢區", "section": section, "subsection": "",
                    "land_number": land_number, "owner_name": owner, "name": owner}

        return {
            "source_file_name": "地主清冊.xlsx",
            "column_map": {1: "district", 2: "section", 3: "land_number", 4: "owner_name"},
            "records": [
                record("龍岡段", "221-0000", "新甲"),  # a land that is already in 龍岡
                record("新段", "1-0000", "新乙"),
                record("新段", "2-0000", "新丙"),
            ],
            "error_rows": [],
        }

    def open_dialog(self, window, kind):
        return next(d for d in reversed(window._open_edit_dialogs) if isinstance(d, kind))

    def run_import(self, window, plan_id):
        from customer_dialogs import ImportPreviewDialog, ImportResultDialog

        window.handle_excel_import_ready(self.import_result())
        preview = self.open_dialog(window, ImportPreviewDialog)
        if plan_id is not None:
            preview.plan_choice.combo.setCurrentIndex(preview.plan_choice.combo.findData(plan_id))
        preview.import_all()
        self.settle(window)
        return preview, self.open_dialog(window, ImportResultDialog)

    def test_the_import_preview_offers_the_plans_and_defaults_to_input_plan(self):
        from customer_dialogs import ImportPreviewDialog

        server, _repository, window = self.open_window()
        window.plan_input_combo.setCurrentIndex(window.plan_input_combo.findData(2))
        window.handle_excel_import_ready(self.import_result())
        preview = self.open_dialog(window, ImportPreviewDialog)
        self.assertEqual(preview.plan_id(), 2)
        self.assertEqual(
            [preview.plan_choice.combo.itemText(i) for i in range(preview.plan_choice.combo.count())],
            ["不指定（未分類）", "龍岡都市計畫", "大湳都市計畫", "空的都市計畫"],
        )
        preview.reject()

    def test_importing_into_a_plan_never_moves_lands_that_already_have_one(self):
        server, _repository, window = self.open_window()
        _preview, result = self.run_import(window, 2)
        self.assertEqual(server.imports[-1], (3, 2))
        plans_by_land = {row["land_id"]: row["urban_plan_name"] for row in server.rows}
        self.assertEqual(plans_by_land[10], "龍岡都市計畫")  # not overwritten
        new_lands = [row for row in server.rows if row["owner_name"] in ("新乙", "新丙")]
        self.assertEqual({row["urban_plan_name"] for row in new_lands}, {"大湳都市計畫"})
        section = result.plan_section
        self.assertIsNotNone(section)
        self.assertIn("2 筆土地新歸入", section.header_label.text())
        self.assertIn("大湳都市計畫", section.header_label.text())
        self.assertEqual(section.table.rowCount(), 1)
        self.assertEqual(
            [section.table.item(0, column).text() for column in range(3)],
            ["中壢區龍岡段", "221-0000", "龍岡都市計畫"],
        )
        self.assertIn("1 筆土地原本已屬於其他都市計畫", section.kept_label.text())

    def test_the_kept_lands_list_can_be_exported(self):
        from openpyxl import load_workbook

        server, _repository, window = self.open_window()
        _preview, result = self.run_import(window, 2)
        target = str(self.root / "kept.xlsx")
        with patch.object(plan_dialogs.QFileDialog, "getSaveFileName", return_value=(target, "")):
            self.assertTrue(result.plan_section.export_kept_lands())
        sheet = load_workbook(target).active
        self.assertEqual(
            [[cell.value for cell in row] for row in sheet.iter_rows()],
            [["地段", "地號", "目前所屬都市計畫"], ["中壢區龍岡段", "221-0000", "龍岡都市計畫"]],
        )

    def test_an_import_without_a_chosen_plan_looks_exactly_like_before(self):
        server, _repository, window = self.open_window()
        _preview, result = self.run_import(window, 0)
        self.assertEqual(server.imports[-1], (3, None))
        self.assertIsNone(result.plan_section)
        self.assertTrue(all(row["urban_plan_id"] is None for row in server.rows if row["land_id"] > 900))

    def test_an_import_with_nothing_kept_has_no_list_or_export_button(self):
        from customer_urban_plan_dialogs import ImportPlanResultSection

        section = ImportPlanResultSection(
            {"name": "龍岡都市計畫", "assigned_land_count": 4, "kept_land_count": 0, "kept_lands": []}
        )
        self.assertIsNone(section.table)
        self.assertIsNone(section.export_button)
        capped = ImportPlanResultSection({
            "name": "龍岡都市計畫", "assigned_land_count": 0, "kept_land_count": 600,
            "kept_lands": [{"district": "中壢區", "section": "龍岡段", "land_number": str(n),
                            "urban_plan_name": "大湳都市計畫"} for n in range(500)],
        })
        self.assertEqual(capped.table.rowCount(), 500)
        self.assertIn("600 筆土地", capped.kept_label.text())

    def test_the_old_preview_dialog_signature_still_works(self):
        from customer_dialogs import ImportPreviewDialog

        dialog = ImportPreviewDialog([], [("district", "地區")], set())
        self.assertEqual(dialog.plan_id(), 0)
        self.assertIsNone(dialog.plan_choice)

    def test_batch_adding_owners_to_one_parcel_uses_the_form_plan(self):
        from customer_dialogs import ImportPreviewDialog

        server, _repository, window = self.open_window()
        window.plan_input_combo.setCurrentIndex(window.plan_input_combo.findData(2))
        window.new_record()
        base_data = window.get_form_data()
        self.assertEqual(base_data["urban_plan_id"], 2)
        shared = {"district": "中壢區", "section": "批新段", "land_number": "9-0000"}
        window._continue_batch_add_shared_land(
            base_data, shared, [{"owner_name": "批甲", "registration_order": "0001"}]
        )
        self.open_dialog(window, ImportPreviewDialog).import_all()
        self.settle(window)
        self.assertEqual(server.imports[-1], (1, 2))
        batch_row = next(row for row in server.rows if row["owner_name"] == "批甲")
        self.assertEqual(batch_row["urban_plan_name"], "大湳都市計畫")

    def test_viewers_cannot_open_a_working_management_dialog(self):
        server, repository, window = self.open_window()
        window.current_user = {**window.current_user, "role": "viewer"}
        window.manage_urban_plans()
        dialog = window.urban_plan_management_dialog
        self.assertFalse(dialog.add_button.isEnabled())
        self.assertFalse(dialog.name_input.isEnabled())
        dialog.close()


if __name__ == "__main__":
    unittest.main()
